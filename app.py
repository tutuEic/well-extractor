"""
地下水统测信息提取器 Web 应用
"""
import sys, os, shutil, uuid, json, re
from pathlib import Path
from datetime import datetime
from fastapi import FastAPI, UploadFile, File
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
import uvicorn
import pytesseract
from PIL import Image

pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

BASE_DIR = Path(__file__).parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "output"
UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

TEMPLATE_PATH = r"D:\地下水\大同朔州忻州_地下水统测表_补全_最终版.xlsx"

app = FastAPI(title="地下水统测信息提取器")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])

sessions = {}

# ===== V4 Extraction Logic (inlined) =====
GW_介质 = ["孔隙水", "岩溶水", "裂隙水", "孔隙-裂隙水"]
GW_深度 = ["浅层", "中层", "中深层", "深层"]
GW_承压 = ["潜水", "承压水", "微承压水"]

def fuzzy_gw(text, candidates, min_ratio=0.4):
    text = text.replace(" ", "").replace("\n", "")
    best, best_score = None, 0
    for c in candidates:
        score = sum(1 for ch in c if ch in text)
        if score > best_score and score >= len(c) * min_ratio:
            best_score, best = score, c
    return best

def extract_single(img_path):
    img = Image.open(img_path)
    full = pytesseract.image_to_string(img, lang="chi_sim+eng", config="--psm 3")
    clean = re.sub(r'(?<=[\u4e00-\u9fff])\s+(?=[\u4e00-\u9fff])', '', full)
    lines = clean.split("\n")
    results = {}
    
    def try_get(patterns, key, post=None, search_text=None):
        text = search_text or clean
        for p in patterns:
            m = re.search(p, text)
            if m:
                v = m.group(1).strip()
                if v in ('/', '无', '请选择', '请输入', '请选择审核人'): continue
                if post: v = post(v)
                if v: results[key] = v
                return
    
    # Well type from title bar
    for line in lines[:5]:
        m = re.search(r'统\s*[测调].*[(（]([^)）]+)[)）]', line)
        if m:
            wt = re.sub(r'[国日\"\'\\]+', '', m.group(1).strip())
            if wt and len(wt) <= 10: results["井点类型"] = wt; break
    
    # Route code
    for line in lines[:5]:
        m = re.match(r'^([A-Za-z]+[-]\d+)\s*$', line.strip())
        if m: results["路线编号"] = m.group(1); break
    
    # IDs
    try_get([r'[野量]外编号[：:\s]*(\S+)', r'编号[：:\s]*([A-Z][A-Z0-9]+)'], "野外编号")
    try_get([r'统[测调][日查]期[：:\s]*(\S+)'], "日期")
    try_get([r'天气[：:\s]*(\S+)'], "天气", lambda v: re.sub(r'[一园国日\"\'\\]+', '', v).strip())
    try_get([r'地理位置[：:]?\s*(.{8,80})'], "地理位置")
    
    # Numeric
    for field, pats in [
        ("井台高度", [r'井台高度[^0-9]*(\d+\.?\d*)']),
        ("井深", [r'[些井此]深[^0-9]*(\d+\.?\d*)']),
        ("测点距地面高度", [r'距地面高度[^0-9]*(\d+\.?\d*)']),
        ("地下水位埋深", [r'[地下]*[水位]*埋深[^0-9]*(\d+\.?\d*)']),
        ("测点距水面距离", [r'距水面距离[^0-9]*(\d+\.?\d*)']),
    ]:
        try_get(pats, field, lambda v: str(float(v)) if v and float(v) < 10000 else None)
    
    # Classifications
    try_get([r'所属统测区类型[：:\s]*(\S+)', r'统测区类型[：:\s]*(\S+)'], "所属统测区类型",
            lambda v: re.sub(r'[航国\"\'\\]+', '', v).strip())
    try_get([r'统测期[次]?[：:\s]*(\S+)'], "统测期", lambda v: re.sub(r'[eEoO0\"\'\\一©国]+', '', v).strip())
    try_get([r'统测[区期]次类型[：:\s]*(\S+)', r'期次类型[：:\s]*(\S+)'], "统测期次类型",
            lambda v: re.sub(r'[国团\"\'\\一&]+', '', v).strip())
    
    # Groundwater types
    for i, line in enumerate(lines):
        if '含水介' in line or '含水层' in line or '承压' in line:
            context = line
            for j in range(i+1, min(i+3, len(lines))): context += " " + lines[j]
            context = context.replace(" ", "")
            if '含水介' in context:
                gw = fuzzy_gw(context, GW_介质, min_ratio=0.3)
                if gw: results["含水介质"] = gw
            if '含水层' in context or '埋藏' in context:
                gw = fuzzy_gw(context, GW_深度)
                if gw: results["埋藏深度"] = gw
            if '承压' in context:
                gw = fuzzy_gw(context, GW_承压)
                if gw: results["承压性"] = gw
    
    # Personnel
    try_get([r'测量人[：:\s]*(\S+)'], "测量人", lambda v: re.sub(r'[国园\"\'\\~\-一¥]+', '', v).strip())
    try_get([r'记录人[：:\s]*(\S+)'], "记录人", lambda v: re.sub(r'[国园\"\'\\~¥]+', '', v).strip())
    try_get([r'审核人[：:\s]*(\S+)'], "审核人",
            lambda v: v if v not in ('请选择审核人',) and not re.search(r'[盆地|山区|平剖]', v) else None)
    try_get([r'实施单位[：:\s]*(.{4,60})', r'统测实施单位[：:\s]*(.{4,60})'], "统测实施单位",
            lambda v: '山西省地质调查院有限公司' if '山西' in v or 'sa' in v.lower() else v)
    try_get([r'项目名称[：:\s]*(.{4,40})'], "项目名称",
            lambda v: v.split('日')[0].split('口')[0].split('山西')[0].strip() or '山西省地下水水位统测')
    
    return results

def post_process(r):
    if not r.get("所属统测区类型") or r["所属统测区类型"] in ("一统测区", ""):
        r["所属统测区类型"] = "一般统测区"
    if not r.get("统测期") or r["统测期"] in ('"', "”", "期", "类型", "") or len(str(r.get("统测期", ""))) > 10:
        r["统测期"] = "地下水平水位期"
    if not r.get("统测期次类型") or r["统测期次类型"] in ("期统测", ""):
        r["统测期次类型"] = "一期统测"
    if not r.get("统测实施单位"): r["统测实施单位"] = "山西省地质调查院有限公司"
    if not r.get("天气"): r["天气"] = "阴"
    if r.get("测量人") and "杨光" in str(r["测量人"]): r["测量人"] = "杨光"
    if r.get("记录人") and "杨升伟" in str(r["记录人"]): r["记录人"] = "杨升伟"


# ===== API Routes =====

@app.get("/")
async def index():
    return HTMLResponse(open(BASE_DIR / "static" / "index.html", encoding="utf-8").read())

@app.post("/api/upload")
async def upload(files: list[UploadFile] = File(...)):
    session_id = str(uuid.uuid4())[:8]
    session_dir = UPLOAD_DIR / session_id
    session_dir.mkdir(exist_ok=True)
    
    records = []
    for f in files:
        if not f.filename: continue
        file_path = session_dir / f.filename
        with open(file_path, "wb") as buf:
            shutil.copyfileobj(f.file, buf)
        try:
            record = extract_single(str(file_path))
            record["_filename"] = f.filename
            records.append(record)
        except Exception as e:
            records.append({"_filename": f.filename, "_error": str(e), "野外编号": "ERROR"})
    
    sessions[session_id] = {"records": records, "created": datetime.now().isoformat()}
    for r in records: post_process(r)
    
    return {"session_id": session_id, "count": len(records), "records": records}

@app.post("/api/session/{session_id}/update")
async def update_record(session_id: str, data: dict):
    if session_id not in sessions: return JSONResponse({"error": "Not found"}, 404)
    idx, field, value = data.get("index"), data.get("field"), data.get("value")
    if idx is None or field is None: return JSONResponse({"error": "Missing field"}, 400)
    records = sessions[session_id]["records"]
    if idx < len(records): records[idx][field] = value; return {"ok": True}
    return JSONResponse({"error": "Invalid index"}, 400)

@app.post("/api/session/{session_id}/export")
async def export_excel(session_id: str):
    if session_id not in sessions: return JSONResponse({"error": "Not found"}, 404)
    output_path = OUTPUT_DIR / f"filled_{session_id}.xlsx"
    try:
        _fill_excel(sessions[session_id]["records"], str(output_path))
        return FileResponse(str(output_path), filename=f"地下水统测表_{session_id}.xlsx",
                          media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        return JSONResponse({"error": str(e)}, 500)

def _fill_excel(records, output_path):
    import openpyxl
    col_map = {"野外编号":"C","统一编号":"B","路线编号":"D","日期":"E","地理位置":"F","天气":"G",
               "井台高度":"P","井深":"Q","测点距地面高度":"U","地下水位埋深":"V","测点距水面距离":"W",
               "井点类型":"T","所属统测区类型":"S","统测期":"AB","统测期次类型":"AC",
               "含水介质":"Y","埋藏深度":"Z","承压性":"AA",
               "测量人":"AH","记录人":"AI","审核人":"AJ","统测实施单位":"AG","项目名称":"AF"}
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb["Sheet1"]
    updated = 0
    for record in records:
        fid = record.get("野外编号")
        if not fid: continue
        target = None
        for row in range(2, ws.max_row + 1):
            if ws[f"C{row}"].value and str(ws[f"C{row}"].value).strip() == str(fid).strip():
                target = row; break
        if target is None: target = ws.max_row + 1
        for key, col in col_map.items():
            if key in record and record[key]:
                try:
                    if key in ("井台高度","井深","测点距地面高度","地下水位埋深","测点距水面距离"):
                        ws[f"{col}{target}"].value = float(record[key])
                    else:
                        ws[f"{col}{target}"].value = str(record[key])
                except: pass
        updated += 1
    wb.save(output_path)

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=8100)
