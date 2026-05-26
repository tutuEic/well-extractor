import openpyxl

wb = openpyxl.load_workbook('D:/地下水/E2E_v2_IDMATCH.xlsx', data_only=False)
ws = wb['Sheet1']

print('Row | Field ID     | 地面高程(L) | 埋深(V) | 水位标高(X) | 测点高程(N)')
print('-' * 80)
start = max(2, ws.max_row - 12)
for row in range(start, ws.max_row + 1):
    fid = ws.cell(row=row, column=3).value
    L = ws.cell(row=row, column=12).value   # 地面高程
    V = ws.cell(row=row, column=22).value   # 地下水位埋深
    X = ws.cell(row=row, column=24).value   # 水位标高
    N = ws.cell(row=row, column=14).value   # 测点高程
    print('{:4d} | {:12s} | {:11s} | {:8s} | {:11s} | {:11s}'.format(
        row, str(fid) if fid else '?',
        str(L) if L else 'EMPTY',
        str(V) if V else 'EMPTY',
        str(X) if X else 'EMPTY',
        str(N) if N else 'EMPTY'))

wb.close()
