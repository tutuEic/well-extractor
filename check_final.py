import openpyxl

wb = openpyxl.load_workbook('D:/地下水/E2E_v2_IDMATCH.xlsx', data_only=False)
ws = wb['Sheet1']

print('Row | ID          | 地面高程 | 埋深  | 水位标高 | 测点高程 | 地理位置')
print('-' * 95)
start = max(2, ws.max_row - 12)
for row in range(start, ws.max_row + 1):
    fid = ws.cell(row=row, column=3).value
    L = ws.cell(row=row, column=12).value
    V = ws.cell(row=row, column=22).value
    X = ws.cell(row=row, column=24).value
    N = ws.cell(row=row, column=14).value
    F = ws.cell(row=row, column=6).value
    x_str = str(X)[:20] if X else 'EMPTY'
    n_str = str(N)[:12] if N else 'EMPTY'
    f_str = str(F)[:35] if F else 'EMPTY'
    print('{:4d} | {:11s} | {:8s} | {:5s} | {:20s} | {:12s} | {}'.format(
        row, str(fid) if fid else '?',
        str(L)[:8] if L else 'EMPTY',
        str(V)[:5] if V else 'EMP',
        x_str, n_str, f_str))
wb.close()
