import openpyxl
wb = openpyxl.load_workbook('D:/地下水/E2E_v2_FULL.xlsx', data_only=False)
ws = wb['Sheet1']
print('Row | ID          | 含水介质 | 埋藏深度 | 承压性')
print('-' * 60)
for row in range(max(2, ws.max_row-12), ws.max_row+1):
    fid = ws.cell(row=row, column=3).value
    Y = ws.cell(row=row, column=25).value
    Z = ws.cell(row=row, column=26).value
    AA = ws.cell(row=row, column=27).value
    print('{:4d} | {:11s} | {:8s} | {:8s} | {:8s}'.format(
        row, str(fid) if fid else '?',
        str(Y) if Y else 'EMPTY',
        str(Z) if Z else 'EMPTY',
        str(AA) if AA else 'EMPTY'))
wb.close()
