import openpyxl
wb = openpyxl.load_workbook('D:/地下水/E2E_COMPLETE.xlsx', data_only=False)
ws = wb['Sheet1']
print('Row | ID          | 地理位置')
print('-' * 55)
for row in range(max(2, ws.max_row-12), ws.max_row+1):
    fid = ws.cell(row=row, column=3).value
    F = ws.cell(row=row, column=6).value
    print('{:4d} | {:11s} | {}'.format(
        row, str(fid) if fid else '?',
        str(F)[:45] if F else 'EMPTY'))
wb.close()
