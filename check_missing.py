import openpyxl

wb = openpyxl.load_workbook('D:/地下水/E2E_v2_IDMATCH.xlsx', data_only=False)
ws = wb['Sheet1']

print('Row | ID          | 含水介质(Y) | 埋藏深度(Z) | 承压性(AA) | 地理位置')
print('-' * 90)
start = max(2, ws.max_row - 12)
for row in range(start, ws.max_row + 1):
    fid = ws.cell(row=row, column=3).value
    Y = ws.cell(row=row, column=25).value   # 含水介质
    Z = ws.cell(row=row, column=26).value   # 埋藏深度
    AA = ws.cell(row=row, column=27).value  # 承压性
    F = ws.cell(row=row, column=6).value    # 地理位置
    print('{:4d} | {:11s} | {:11s} | {:11s} | {:11s} | {}'.format(
        row, str(fid) if fid else '?',
        str(Y) if Y else 'EMPTY',
        str(Z) if Z else 'EMPTY',
        str(AA) if AA else 'EMPTY',
        str(F)[:40] if F else 'EMPTY'))
wb.close()
