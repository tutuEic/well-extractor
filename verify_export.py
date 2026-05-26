import openpyxl

wb = openpyxl.load_workbook('D:/地下水/E2E_v2_FINAL.xlsx', data_only=False)
ws = wb['Sheet1']

print('Verifying E2E_v2_FINAL.xlsx')
print('Row | Field ID     | Longitude           | X               | Elevation   | Depth')
print('-' * 90)

start = max(2, ws.max_row - 12)
for row in range(start, ws.max_row + 1):
    fid = ws.cell(row=row, column=3).value
    lng = ws.cell(row=row, column=8).value
    x = ws.cell(row=row, column=10).value
    elev = ws.cell(row=row, column=12).value
    depth = ws.cell(row=row, column=17).value
    lng_s = str(lng)[:28] if lng else '?'
    x_s = str(x)[:15] if x else '?'
    print('{:4d} | {:12s} | {:28s} | {:15s} | {:10s} | {}'.format(
        row, str(fid) if fid else '?', lng_s, x_s, str(elev) if elev else '?', depth))

wb.close()
