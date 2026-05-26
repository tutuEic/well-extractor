import openpyxl

wb = openpyxl.load_workbook('D:/地下水/E2E_v2_IDMATCH.xlsx', data_only=False)
ws = wb['Sheet1']

print('Row | Field ID     | 埋深(V) | 水面距(W) | 井深(Q) | 地面高(U) | 井台高(P)')
print('-' * 85)
start = max(2, ws.max_row - 12)
for row in range(start, ws.max_row + 1):
    fid = ws.cell(row=row, column=3).value
    v = ws.cell(row=row, column=22).value   # 地下水位埋深
    w = ws.cell(row=row, column=23).value   # 测点距水面距离
    q = ws.cell(row=row, column=17).value   # 井深
    u = ws.cell(row=row, column=21).value   # 测点距地面高度
    p = ws.cell(row=row, column=16).value   # 井台高度
    print('{:4d} | {:12s} | {:8s} | {:8s} | {:8s} | {:8s} | {:8s}'.format(
        row, str(fid) if fid else '?',
        str(v) if v else 'EMPTY',
        str(w) if w else 'EMPTY',
        str(q) if q else 'EMPTY',
        str(u) if u else 'EMPTY',
        str(p) if p else 'EMPTY'))

wb.close()
